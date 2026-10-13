"""Master Excel workbook: load, merge without duplicates, save.

The workbook is the source of truth. You can edit it by hand (e.g. set a
Status to Done, fix a due date, add Notes); manual edits are preserved.
"""
from __future__ import annotations

import re
import shutil
from datetime import date, datetime, timedelta
from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import settings
from .models import Extraction
from .sources import SourceItem

SHEETS: dict[str, list[str]] = {
    "Commitments": [
        "ID", "Owner", "Owner Email", "Item", "Due", "Status", "Project",
        "Source", "Source Link", "Evidence", "First Seen", "Last Updated", "Done On", "Notes",
    ],
    "Decisions": [
        "ID", "Decision", "Made By", "Date", "Project",
        "Source", "Source Link", "Evidence", "First Seen", "Notes",
    ],
    "Projects": [
        "Project", "Open Commitments", "Overdue", "Pending Approvals", "Open Risks",
        "High Risks", "Decisions (30d)", "Last Activity", "Notes",
    ],
    "People": [
        "Name", "Email", "Phone", "Organisation", "Role", "First Seen", "Last Seen",
        "Open Commitments", "Notes",
    ],
    "Approvals": [
        "ID", "Item", "Requested By", "Requested By Email", "Amount", "Due", "Status", "Project",
        "Source", "Source Link", "Evidence", "First Seen", "Last Updated", "Resolved On", "Notes",
    ],
    # Not in the original brief but needed to keep extracted risks de-duplicated over time.
    "Risks": [
        "ID", "Risk", "Severity", "Project", "Owner", "Status",
        "Source", "Source Link", "Evidence", "First Seen", "Last Seen", "Notes",
    ],
}
ID_PREFIX = {"Commitments": "C", "Decisions": "D", "Approvals": "A", "Risks": "R"}

OPEN_COMMITMENT = {"Open", "Blocked"}
OPEN_APPROVAL = {"Pending"}
SIMILAR = 0.86


def norm(s) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(s or "").lower()).strip()


def similar(a, b) -> bool:
    a, b = norm(a), norm(b)
    if not a or not b:
        return False
    return a == b or SequenceMatcher(None, a, b).ratio() >= SIMILAR


def _cell(v):
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, date):
        return v.isoformat()
    return v


class Store:
    def __init__(self, path: Path | None = None):
        self.path = path or settings.master_xlsx
        self.rows: dict[str, list[dict]] = {name: [] for name in SHEETS}
        if self.path.exists():
            self._load()

    # ---------- IO ----------
    def _load(self) -> None:
        wb = load_workbook(self.path)
        for name in SHEETS:
            if name not in wb.sheetnames:
                continue
            ws = wb[name]
            values = list(ws.iter_rows(values_only=True))
            if not values:
                continue
            header = [str(h) if h is not None else "" for h in values[0]]
            for r in values[1:]:
                if not any(v not in (None, "") for v in r):
                    continue
                self.rows[name].append({h: _cell(v) for h, v in zip(header, r) if h})

    def save(self) -> Path:
        wb = Workbook()
        wb.remove(wb.active)
        head_font = Font(bold=True, color="FFFFFF")
        head_fill = PatternFill("solid", fgColor="1F3A5F")
        for name, base_cols in SHEETS.items():
            ws = wb.create_sheet(name)
            # keep any extra columns the user added by hand
            extra = []
            for row in self.rows[name]:
                extra += [k for k in row if k not in base_cols and k not in extra]
            cols = base_cols + extra
            ws.append(cols)
            for row in self.rows[name]:
                ws.append([row.get(c) for c in cols])
            for i, c in enumerate(cols, 1):
                cell = ws.cell(row=1, column=i)
                cell.font, cell.fill = head_font, head_fill
                wide = c in {"Item", "Decision", "Risk", "Evidence", "Notes"}
                ws.column_dimensions[get_column_letter(i)].width = 60 if wide else 18
            for r in ws.iter_rows(min_row=2):
                for cell in r:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.freeze_panes = "B2"
            ws.auto_filter.ref = ws.dimensions

        self.path.parent.mkdir(parents=True, exist_ok=True)
        tmp = self.path.with_suffix(".tmp.xlsx")
        wb.save(tmp)
        try:
            if self.path.exists():
                shutil.copy2(self.path, self.path.with_suffix(".bak.xlsx"))
            tmp.replace(self.path)
            return self.path
        except PermissionError:
            alt = self.path.with_name(
                f"{self.path.stem}.{datetime.now():%Y%m%d-%H%M}.xlsx"
            )
            tmp.replace(alt)
            print(f"[excel] {self.path.name} is open in Excel; wrote {alt.name} instead. "
                  "Close Excel and rename it to keep one master file.")
            return alt

    # ---------- helpers ----------
    def _next_id(self, sheet: str) -> str:
        prefix = ID_PREFIX[sheet]
        nums = [
            int(m.group(1))
            for r in self.rows[sheet]
            if (m := re.match(rf"{prefix}-(\d+)$", str(r.get("ID", ""))))
        ]
        return f"{prefix}-{(max(nums) + 1 if nums else 1):04d}"

    def _by_id(self, sheet: str, id_: str) -> dict | None:
        return next((r for r in self.rows[sheet] if r.get("ID") == id_), None)

    def open_ledger(self) -> dict:
        """Compact view of open items for Claude to match against."""
        return {
            "open_commitments": [
                {k: r.get(k) for k in ("ID", "Owner", "Item", "Due", "Status", "Project")}
                for r in self.rows["Commitments"] if r.get("Status") in OPEN_COMMITMENT
            ],
            "pending_approvals": [
                {k: r.get(k) for k in ("ID", "Item", "Requested By", "Due", "Project")}
                for r in self.rows["Approvals"] if r.get("Status") in OPEN_APPROVAL
            ],
            "known_projects": sorted({r["Project"] for r in self.rows["Projects"] if r.get("Project")}),
        }

    # ---------- merge ----------
    def merge(self, ex: Extraction, items: dict[str, SourceItem], now: datetime) -> dict:
        stamp = now.strftime("%Y-%m-%d %H:%M")
        stats = {k: 0 for k in ("commitments", "closed", "decisions", "risks", "approvals", "resolved", "people")}

        def src(ref: str) -> tuple[str, str]:
            it = items.get(ref)
            return (f"{it.source}: {it.title}", it.link) if it else (ref, "")

        # Commitments: updates to existing items first
        for u in ex.commitment_updates:
            row = self._by_id("Commitments", u.id)
            if not row:
                continue
            if u.status != row.get("Status"):
                row["Status"] = u.status
                if u.status == "Done":
                    row["Done On"] = now.date().isoformat()
                    stats["closed"] += 1
            if u.due:
                row["Due"] = u.due
            row["Evidence"] = f"{row.get('Evidence') or ''}\n[{stamp}] {u.evidence}".strip()
            row["Last Updated"] = stamp

        for c in ex.commitments:
            dup = next(
                (r for r in self.rows["Commitments"]
                 if similar(r.get("Owner"), c.owner) and similar(r.get("Item"), c.item)),
                None,
            )
            if dup:
                if c.status == "Done" and dup.get("Status") in OPEN_COMMITMENT:
                    dup["Status"], dup["Done On"] = "Done", now.date().isoformat()
                    dup["Last Updated"] = stamp
                    stats["closed"] += 1
                if c.due and not dup.get("Due"):
                    dup["Due"] = c.due
                continue
            s, link = src(c.source_ref)
            self.rows["Commitments"].append({
                "ID": self._next_id("Commitments"), "Owner": c.owner, "Owner Email": c.owner_email,
                "Item": c.item, "Due": c.due, "Status": c.status, "Project": c.project,
                "Source": s, "Source Link": link, "Evidence": c.evidence,
                "First Seen": stamp, "Last Updated": stamp,
                "Done On": now.date().isoformat() if c.status == "Done" else None,
            })
            stats["commitments"] += 1

        for d in ex.decisions:
            if any(similar(r.get("Decision"), d.decision) for r in self.rows["Decisions"]):
                continue
            s, link = src(d.source_ref)
            self.rows["Decisions"].append({
                "ID": self._next_id("Decisions"), "Decision": d.decision, "Made By": d.made_by,
                "Date": d.date, "Project": d.project, "Source": s, "Source Link": link,
                "Evidence": d.evidence, "First Seen": stamp,
            })
            stats["decisions"] += 1

        for rk in ex.risks:
            dup = next((r for r in self.rows["Risks"] if similar(r.get("Risk"), rk.risk)), None)
            if dup:
                dup["Last Seen"] = stamp
                if rk.severity == "High":
                    dup["Severity"] = "High"
                continue
            s, link = src(rk.source_ref)
            self.rows["Risks"].append({
                "ID": self._next_id("Risks"), "Risk": rk.risk, "Severity": rk.severity,
                "Project": rk.project, "Owner": rk.owner, "Status": "Open",
                "Source": s, "Source Link": link, "Evidence": rk.evidence,
                "First Seen": stamp, "Last Seen": stamp,
            })
            stats["risks"] += 1

        for u in ex.approval_updates:
            row = self._by_id("Approvals", u.id)
            if not row or u.status == row.get("Status"):
                continue
            row["Status"] = u.status
            row["Evidence"] = f"{row.get('Evidence') or ''}\n[{stamp}] {u.evidence}".strip()
            row["Last Updated"] = stamp
            if u.status != "Pending":
                row["Resolved On"] = now.date().isoformat()
                stats["resolved"] += 1

        for a in ex.approvals:
            if any(similar(r.get("Item"), a.item) for r in self.rows["Approvals"]):
                continue
            s, link = src(a.source_ref)
            self.rows["Approvals"].append({
                "ID": self._next_id("Approvals"), "Item": a.item, "Requested By": a.requested_by,
                "Requested By Email": a.requested_by_email, "Amount": a.amount, "Due": a.due,
                "Status": "Pending", "Project": a.project, "Source": s, "Source Link": link,
                "Evidence": a.evidence, "First Seen": stamp, "Last Updated": stamp,
            })
            stats["approvals"] += 1

        for p in ex.people:
            if p.email and p.email.lower() in settings.me_emails:
                continue
            dup = next(
                (r for r in self.rows["People"]
                 if (p.email and str(r.get("Email") or "").lower() == p.email.lower())
                 or norm(r.get("Name")) == norm(p.name)),
                None,
            )
            if dup:
                dup["Last Seen"] = stamp
                for col, val in (("Email", p.email), ("Phone", p.phone), ("Organisation", p.organisation), ("Role", p.role)):
                    if val and not dup.get(col):
                        dup[col] = val
                continue
            self.rows["People"].append({
                "Name": p.name, "Email": p.email, "Phone": p.phone, "Organisation": p.organisation,
                "Role": p.role, "First Seen": stamp, "Last Seen": stamp,
            })
            stats["people"] += 1
        return stats

    def refresh_rollups(self, today: date) -> None:
        """Recompute Projects and People counters from the other sheets."""
        open_c = [r for r in self.rows["Commitments"] if r.get("Status") in OPEN_COMMITMENT]
        pending = [r for r in self.rows["Approvals"] if r.get("Status") in OPEN_APPROVAL]
        open_r = [r for r in self.rows["Risks"] if r.get("Status", "Open") == "Open"]
        cutoff = (today - timedelta(days=30)).isoformat()

        names: dict[str, str] = {}
        for sheet in ("Commitments", "Approvals", "Risks", "Decisions", "Projects"):
            for r in self.rows[sheet]:
                if r.get("Project"):
                    names.setdefault(norm(r["Project"]), r["Project"])

        existing = {norm(r["Project"]): r for r in self.rows["Projects"] if r.get("Project")}
        projects = []
        for key, name in sorted(names.items(), key=lambda kv: kv[1].lower()):
            row = existing.get(key, {"Project": name})
            mine = lambda rows: [r for r in rows if norm(r.get("Project")) == key]  # noqa: E731
            pc = mine(open_c)
            activity = [
                str(r.get(col) or "")
                for sheet in ("Commitments", "Approvals", "Risks", "Decisions")
                for r in mine(self.rows[sheet])
                for col in ("Last Updated", "Last Seen", "First Seen")
            ]
            row.update({
                "Open Commitments": len(pc),
                "Overdue": sum(1 for r in pc if r.get("Due") and str(r["Due"]) < today.isoformat()),
                "Pending Approvals": len(mine(pending)),
                "Open Risks": len(mine(open_r)),
                "High Risks": sum(1 for r in mine(open_r) if r.get("Severity") == "High"),
                "Decisions (30d)": sum(1 for r in mine(self.rows["Decisions"]) if str(r.get("Date") or "") >= cutoff),
                "Last Activity": max(activity, default="")[:10] or None,
            })
            projects.append(row)
        self.rows["Projects"] = projects

        for p in self.rows["People"]:
            p["Open Commitments"] = sum(
                1 for r in open_c
                if norm(r.get("Owner")) == norm(p.get("Name"))
                or (p.get("Email") and str(r.get("Owner Email") or "").lower() == str(p["Email"]).lower())
            )
